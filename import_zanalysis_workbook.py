#!/usr/bin/env python3
"""
Import zAnalysis tables kept in Excel workbooks (e.g. Zachary Diaz's
labnotes/*.xlsx) as per-folder zAnalysis configs under data/<date>/.

Each workbook's first sheet holds a zAnalysis table in columns A-N (header
in row 1); anything to the right (summary tables, deviations) is ignored.
Rows are split by the date in the Gamry filename (..._YYYYMMDD_HHMMSS_...),
so a run spanning midnight lands in both day folders. Every row gets a note
naming the source workbook; a day without its own standard row gets a note
naming the day whose standard set its conductivities. Row font colours that
the workbook explains in a legend (e.g. 'red = before', 'blue = after') are
carried into the notes. The CSV is then harmonized to JSON.

Usage:
    python import_zanalysis_workbook.py labnotes/*.xlsx [--force] [--dry-run] [--note TEXT]
"""
import argparse
import glob
import os
import re

import openpyxl
import pandas as pd

COLUMNS = ['group_name', 'filename', 'P_MPa', 'T_K', 'type', 'Z_Ohm',
           'Z_unc_Ohm', 'conductivity_Sm', 'S_unc_pct', 'comp', 'w_ppt',
           'w_molal', 'exclude', 'notes']
DATE = re.compile(r'_(\d{8})_\d{6}_')
COLOUR_NAMES = {'FFFF0000': 'red', 'FF0070C0': 'blue'}


def colour_legend(ws):
    """Map colour name -> meaning from cells like 'red = before'."""
    legend = {}
    for row in ws.iter_rows(min_col=len(COLUMNS) + 1):
        for c in row:
            m = re.match(r'\s*(\w+)\s*=\s*(.+)', str(c.value or ''))
            if m and m.group(1).lower() in COLOUR_NAMES.values():
                legend[m.group(1).lower()] = m.group(2).strip()
    return legend


def read_workbook(path):
    ws = openpyxl.load_workbook(path).worksheets[0]
    header = [ws.cell(1, j).value for j in range(1, len(COLUMNS) + 1)]
    if header != COLUMNS:
        raise ValueError(f'{path}: unexpected header {header}')
    legend = colour_legend(ws)
    rows = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, j).value for j in range(1, len(COLUMNS) + 1)]
        if not vals[1]:
            continue
        rec = dict(zip(COLUMNS, vals))
        font = ws.cell(r, 8).font
        rgb = font.color.rgb if font and font.color else None
        colour = COLOUR_NAMES.get(rgb) if isinstance(rgb, str) else None
        notes = [f'Z. Diaz analysis ({os.path.basename(path)})']
        if colour in legend:
            notes.append(f'{legend[colour]} Gamry recalibration')
        if rec['notes']:
            notes.insert(0, str(rec['notes']))
        rec['notes'] = '; '.join(notes)
        rows.append(rec)
    return pd.DataFrame(rows, columns=COLUMNS)


def split_by_day(df):
    df = df.copy()
    df['day'] = df['filename'].str.extract(DATE, expand=False)
    if df['day'].isna().any():
        bad = df.loc[df['day'].isna(), 'filename'].tolist()
        raise ValueError(f'no date in filenames: {bad}')
    std_days = sorted(df.loc[df['type'] == 'standard', 'day'].unique())
    out = {}
    for day, g in df.groupby('day'):
        g = g.drop(columns='day').copy()
        if std_days and day not in std_days:
            g['notes'] = g['notes'] + f'; calibrated with the {std_days[0]} standard'
        out[day] = g
    return out


def missing_files(day, df, root):
    present = {os.path.basename(p) for p in
               glob.glob(os.path.join(root, 'data', day, 'Conductivity*', '*.txt'))}
    return [f for f in df['filename'] if f not in present]


def main():
    p = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    p.add_argument('workbooks', nargs='+')
    p.add_argument('--root', default=os.path.dirname(os.path.abspath(__file__)))
    p.add_argument('--force', action='store_true', help='overwrite existing configs')
    p.add_argument('--dry-run', action='store_true')
    p.add_argument('--note', help='extra note appended to every imported row')
    a = p.parse_args()

    for wb in a.workbooks:
        if os.path.basename(wb).startswith('~$'):
            continue  # Excel lock file
        for day, df in split_by_day(read_workbook(wb)).items():
            if a.note:
                df['notes'] = df['notes'] + f'; {a.note}'
            cfg = os.path.join(a.root, 'data', day, f'zAnalysis{day}.csv')
            missing = missing_files(day, df, a.root)
            status = 'exists, skipped' if os.path.exists(cfg) and not a.force else 'write'
            print(f'{day}: {len(df)} rows ({(df["type"] == "standard").sum()} standard) '
                  f'from {os.path.basename(wb)} -> {status}'
                  + (f'; MISSING data files: {missing}' if missing else ''))
            if a.dry_run or status != 'write' or missing:
                continue
            df.to_csv(cfg, index=False)
            from harmonize_config import harmonize_config
            harmonize_config(__import__('pathlib').Path(cfg), force=True, verbose=False)


if __name__ == '__main__':
    main()
